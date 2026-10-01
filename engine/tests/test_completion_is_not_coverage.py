from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook
from vvs_engine.coverage import completion_checks
from app.exports import to_csv, to_xlsx


def test_incomplete_checks_degrade_valid_coverage_without_mutating_it():
    quality = {'verdict': 'VALID', 'reasons': [], 'accuracy_verified': False}
    original = deepcopy(quality)
    result = completion_checks(quality, [{'review_m': 2.3}],
        {'model': {'result': {'assignments': [{'status': 'unresolved', 'reason': 'unanswered'}]}}},
        {'findings': [{'code': 'ocr_partial'}]})
    assert result['verdict'] == 'DEGRADED'
    assert set(result['reasons']) == {'TENTATIVE_QUANTITIES', 'MODEL_REVIEW_INCOMPLETE', 'OCR_REVIEW_INCOMPLETE'}
    assert quality == original
    assert completion_checks(result, [{'review_m': 2.3}]) == result


def test_invalid_stays_invalid_and_recovered_requests_do_not_imply_incomplete():
    assert completion_checks({'verdict': 'INVALID'}, [{'review_m': 1}])['verdict'] == 'INVALID'
    quality = {'verdict': 'VALID', 'reasons': []}
    assert completion_checks(quality, source_assignment={'model': {'result': {
        'transport_failures': [{'error': 'timeout'}], 'assignments': [
            {'status': 'unresolved', 'reason': 'no_candidate'},
            {'status': 'unresolved', 'reason': 'abstained'}]}}}) == quality


def test_exports_identify_tentative_metres():
    row = dict(designation='S3-R8', dn=75, physical_pipe_count=1,
               confirmed_horizontal_m=10, vertical_m='UNKNOWN', confirmed_total_m=10,
               ambiguous_m=0, state='CONFIRMED', review_m=2.75)
    csv = to_csv('', rows=[row])
    assert 'Varav behöver granskas m' in csv and '2.75' in csv
    sheet = load_workbook(BytesIO(to_xlsx('', rows=[row]))).active
    assert sheet.cell(1, 16).value == 'Varav behöver granskas m'
    assert sheet.cell(2, 16).value == 2.75
    assert row['review_m'] == 2.75
